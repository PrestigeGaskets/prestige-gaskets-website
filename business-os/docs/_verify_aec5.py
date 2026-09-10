from pathlib import Path

from openpyxl import load_workbook

p = Path(r"C:\Users\PC\rushmore-voxel-cone\business-os\SMB-Operating-System.xlsx")
wb = load_workbook(p)
print("SHEETS", wb.sheetnames)
print("SIZE", p.stat().st_size)
print("DEFINED", list(wb.defined_names.keys()))
print("--- TABLES ---")
for name in wb.sheetnames:
    ws = wb[name]
    tables = list(ws.tables.keys())
    dvs = len(ws.data_validations.dataValidation)
    print(f"{name:16} tables={tables} dv={dvs} hidden={ws.sheet_state}")

ws = wb["QuoteLines"]
total_by_q = {}
for row in ws.iter_rows(min_row=2, max_row=10, max_col=6, values_only=True):
    qno, line, sku, qty, price, formula = row
    if not qno:
        continue
    val = qty * price
    total_by_q[qno] = total_by_q.get(qno, 0) + val
    print(f"  {qno} L{line} {sku} {qty} x {price} = {val:.2f} formula={formula}")
print("QUOTE TOTALS", {k: round(v, 2) for k, v in total_by_q.items()}, "grand", round(sum(total_by_q.values()), 2))

print("--- PRODUCTS ---")
ws = wb["Products"]
print([c.value for c in ws[1]])
for row in ws.iter_rows(min_row=2, max_row=7, values_only=True):
    print(
        f"  {row[0]} {row[7]} {row[8]} OH={row[14]} ROP={row[18]} lead={row[20]} "
        f"cost={row[10]} sell={row[11]}"
    )

print("--- DASHBOARD ---")
dash = wb["Dashboard"]
for r in range(1, 9):
    print(f"  {r}: {dash[f'A{r}'].value} | {dash[f'B{r}'].value}")
print("Quotes Customer", wb["Quotes"]["D2"].value)
print("Quotes Value", wb["Quotes"]["G2"].value)
print("Orders ref", wb["Orders"].tables["tblOrders"].ref)
print("C004", wb["Customers"]["A5"].value, wb["Customers"]["G5"].value, wb["Customers"]["J5"].value)
print("Lists", wb["Lists"].sheet_state)
print("P1002 flag formula", wb["Products"]["Z3"].value)
